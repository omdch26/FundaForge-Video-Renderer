"""Read the source of truth: Unit_Index.xlsx -> the unit's CSV.

Never reads the rendered PDFs. The CSV is the source; the PDF is a downstream
artefact of a different pipeline and cropping it was explicitly rejected.
"""
from __future__ import annotations

import csv
import hashlib
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .config import Config


@dataclass
class Slide:
    slide_number: int
    slide_type: str
    layout_module: str
    header_text: str
    body_text: str
    visual_direction: str

    @property
    def names_asset(self) -> str | None:
        """visual_direction is a human instruction unless it names a registered asset.

        Mirrors batch_render_season.py: a value starting with 'none' is prose,
        anything else is an asset stem to resolve against the vector library.
        """
        v = (self.visual_direction or "").strip()
        if not v or v.lower().startswith("none"):
            return None
        return v.split()[0].strip(" ,.")


@dataclass
class Unit:
    unit_id: str
    lane: str            # 'season' | 'blueprint'
    series: str
    tier: str
    kicker: str
    hook_slide1: str
    why_slide2: str
    cta_type: str        # 'plain' | 'keyword'
    keyword: str | None
    ig_status: str
    csv_path: str
    slides: list[Slide] = field(default_factory=list)
    csv_sha256: str = ""

    @property
    def is_blueprint(self) -> bool:
        return self.lane == "blueprint"

    @property
    def trap_slide(self) -> Slide:
        """Slide 9. The most video-native slide in the format, and the one
        compression is most likely to flatten."""
        return self.slides[8]


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_index(cfg: Config) -> dict[str, dict]:
    """All 86 rows of Unit_Index.xlsx, keyed by unit_id."""
    wb = openpyxl.load_workbook(cfg.ro("unit_index"), read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) for h in next(rows)]
    out = {}
    for row in rows:
        if not row or not row[0]:
            continue
        rec = dict(zip(headers, row))
        out[str(rec["unit_id"])] = rec
    wb.close()
    return out


def load_unit(cfg: Config, unit_id: str) -> Unit:
    index = load_index(cfg)
    if unit_id not in index:
        raise KeyError(f"{unit_id} is not in Unit_Index.xlsx")
    rec = index[unit_id]

    csv_path = cfg.system_root / str(rec["csv_path"])
    if not csv_path.exists():
        raise FileNotFoundError(f"{unit_id} -> {csv_path} (from Unit_Index csv_path column)")

    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        slides = [
            Slide(
                slide_number=int(r["slide_number"]),
                slide_type=r["slide_type"],
                layout_module=r["layout_module"],
                header_text=r["header_text"],
                body_text=r["body_text"],
                visual_direction=r.get("visual_direction", ""),
            )
            for r in csv.DictReader(fh)
        ]

    if len(slides) != 10:
        raise ValueError(f"{unit_id}: expected 10 slides, found {len(slides)}")

    lane = "blueprint" if str(rec["lane"]).lower().startswith("blue") else "season"

    return Unit(
        unit_id=unit_id,
        lane=lane,
        series=str(rec["series"]),
        tier=str(rec["tier"]),
        kicker=str(rec["kicker"]),
        hook_slide1=str(rec["hook_slide1"]),
        why_slide2=str(rec["why_slide2"]),
        cta_type=str(rec["cta_type"]),
        keyword=(str(rec["keyword"]) if rec.get("keyword") else None),
        ig_status=str(rec["ig_status"]),
        csv_path=str(rec["csv_path"]),
        slides=slides,
        csv_sha256=_sha256(csv_path),
    )
